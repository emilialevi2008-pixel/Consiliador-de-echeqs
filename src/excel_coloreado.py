import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment


ESTADOS_CORRECTOS = {
    "CONCILIADO",
    "CONCILIADO_ID_ABREVIADO",
}


def generar_excel_coloreado(carpeta):
    carpeta = Path(carpeta)
    archivo_csv = carpeta / "reporte_conciliacion.csv"
    archivo_excel = carpeta / "reporte_conciliacion.xlsx"

    with open(archivo_csv, encoding="utf-8-sig", newline="") as archivo:
        lector = csv.DictReader(archivo)
        columnas = lector.fieldnames
        resultados = list(lector)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Conciliación"

    hoja.append(columnas)

    verde = PatternFill("solid", fgColor="E2F0D9")
    rojo = PatternFill("solid", fgColor="FCE4D6")
    azul = PatternFill("solid", fgColor="1F4E78")

    for celda in hoja[1]:
        celda.fill = azul
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(wrap_text=True)

    for resultado in resultados:
        hoja.append([resultado[columna] for columna in columnas])

        color = (
            verde
            if resultado["estado"] in ESTADOS_CORRECTOS
            else rojo
        )

        for celda in hoja[hoja.max_row]:
            celda.fill = color

    anchos = {
        "A": 20, "B": 20, "C": 18, "D": 18,
        "E": 20, "F": 17, "G": 17, "H": 28,
        "I": 28, "J": 65, "K": 15, "L": 15,
    }

    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    libro.save(archivo_excel)
    print(f"Excel con colores generado: {archivo_excel}")


if __name__ == "__main__":
    generar_excel_coloreado("outputs")