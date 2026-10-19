from openpyxl import load_workbook
import csv
import json
from abc import ABC, abstractmethod
from decimal import Decimal


class FuenteDatos(ABC):
    """Define cómo deben funcionar los lectores de archivos."""

    @abstractmethod
    def cargar(self, ruta):
        pass


class FuenteCSV(FuenteDatos):
    """Lee un archivo CSV y devuelve una lista de diccionarios."""

    def cargar(self, ruta):
        with open(ruta, encoding="utf-8-sig", newline="") as archivo:
            lector = csv.DictReader(archivo, strict=True)

            if not lector.fieldnames:
                raise ValueError("El CSV no tiene encabezados.")

            if len(lector.fieldnames) != len(set(lector.fieldnames)):
                raise ValueError("El CSV tiene encabezados repetidos.")

            datos = list(lector)

        if not datos:
            raise ValueError("El CSV no contiene registros.")

        for fila in datos:
            if None in fila or any(valor is None for valor in fila.values()):
                raise ValueError(
                    "Una fila del CSV no coincide con sus columnas."
                )

        return datos


class FuenteJSON(FuenteDatos):
    """Lee un JSON que contiene una lista de registros."""

    def cargar(self, ruta):
        with open(ruta, encoding="utf-8-sig") as archivo:
            datos = json.load(archivo, parse_float=Decimal)

        if not isinstance(datos, list) or not datos:
            raise ValueError(
                "El JSON debe contener una lista de registros no vacía."
            )

        if not all(isinstance(fila, dict) for fila in datos):
            raise ValueError("Cada registro del JSON debe ser un objeto.")

        return datos
class FuenteExcel(FuenteDatos):
    """Lee la primera hoja de un archivo Excel."""

    def cargar(self, ruta):
        libro = load_workbook(
            ruta,
            read_only=True,
            data_only=True
        )

        try:
            hoja = libro.worksheets[0]
            filas = hoja.iter_rows(values_only=True)
            encabezados = next(filas, None)

            if not encabezados:
                raise ValueError("El Excel está vacío.")

            if any(
                not isinstance(nombre, str) or not nombre.strip()
                for nombre in encabezados
            ):
                raise ValueError(
                    "Cada columna del Excel debe tener un encabezado."
                )

            encabezados = [
                nombre.strip() for nombre in encabezados
            ]

            if len(encabezados) != len(set(encabezados)):
                raise ValueError(
                    "El Excel tiene encabezados repetidos."
                )

            datos = []

            for fila in filas:
                if all(valor is None for valor in fila):
                    continue

                registro = dict(zip(encabezados, fila))
                datos.append(registro)

            if not datos:
                raise ValueError(
                    "El Excel no contiene registros."
                )

            return datos

        finally:
            libro.close()